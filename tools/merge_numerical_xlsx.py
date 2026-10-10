"""把 iPad（ScalePad）导出的数轴/点数实验 xlsx 追加进汇总表，效果与原 PsychoPy 程序相同。

原程序每做完一次，就在 “数轴实验_results.xlsx” / “点数实验_results.xlsx” 里新建一个以完成时间命名的 sheet。
ScalePad 每次导出一个只含这一个 sheet 的文件，文件名形如：
    数轴实验_P001_2026-10-11_14-03-22.xlsx
    点数实验_P001_2026-10-11_14-10-05.xlsx
本脚本按文件名前缀找到对应的汇总表，把 sheet 原样复制进去（同名 sheet 已存在则跳过，可重复运行）。

用法：
    python3 merge_numerical_xlsx.py 导出文件1.xlsx [导出文件2.xlsx ...] [--dest 目标文件夹]
目标文件夹默认是 ~/Desktop/numerical cognition（与原程序相同），也可以指定某位患者的文件夹。
需要 openpyxl（PsychoPy 自带；系统 Python 可用 pip3 install openpyxl 安装）。
"""

from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("缺少 openpyxl，请先执行：pip3 install openpyxl")

WORKBOOKS = {"数轴实验": "数轴实验_results.xlsx", "点数实验": "点数实验_results.xlsx"}


def merge(source: Path, dest_dir: Path) -> str:
    prefix = source.name.split("_", 1)[0]
    if prefix not in WORKBOOKS:
        return f"跳过 {source.name}：文件名不是以“数轴实验_”或“点数实验_”开头"
    target = dest_dir / WORKBOOKS[prefix]
    src_wb = openpyxl.load_workbook(source)
    if target.exists():
        dst_wb = openpyxl.load_workbook(target)
    else:
        dst_wb = openpyxl.Workbook()
        dst_wb.remove(dst_wb.active)
    added = []
    for src_ws in src_wb.worksheets:
        if src_ws.title in dst_wb.sheetnames:
            continue
        dst_ws = dst_wb.create_sheet(title=src_ws.title)
        for row in src_ws.iter_rows(values_only=True):
            dst_ws.append(list(row))
        added.append(src_ws.title)
    if not added:
        return f"{source.name}：{target.name} 里已有同名 sheet，未重复追加"
    dst_wb.save(target)
    return f"{source.name} → {target}（新增 sheet：{', '.join(added)}）"


def main() -> None:
    parser = argparse.ArgumentParser(description="把 ScalePad 导出的数轴/点数实验结果追加进汇总 xlsx")
    parser.add_argument("files", nargs="+", type=Path)
    parser.add_argument("--dest", type=Path, default=Path(os.path.expanduser("~")) / "Desktop" / "numerical cognition")
    args = parser.parse_args()
    args.dest.mkdir(parents=True, exist_ok=True)
    for source in args.files:
        print(merge(source, args.dest))


if __name__ == "__main__":
    main()
